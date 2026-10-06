"""Write the daily Excel of the shipments table, in the layout the company knows.

6.10.2026. Run by tools/ship-daily.sh after ship-brief.js:
    SHIP_DIR=/srv/homestyle/shipments /opt/ship-venv/bin/python -I tools/ship-export.py <out.xlsx>
One sheet per sheet of the original file, the same column headings, real dates,
and one extra column with the alert codes. The table is the source; this file
is a view of it.
"""
import datetime as dt
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLS = [("PO#", "po"), ("SAP", "sap"), ("Project", "project"), ("SUPPLIER", "supplier"), ("מוצר", "product"),
        ("ITEM", "item"), ("FCL", "fcl"), ("R.ETD", "retd"), ("ETD", "etd"), ("ETA", "eta"),
        ("תיק עמילות", "customs"), ("CONTAINER", "container"), ("סטטוס", "status"), ("הערות", "notes"),
        ("סטטוס תשלום", "payStatus"), ("סכום הזמנה", "amount")]
DATES = {"retd", "etd", "eta"}


def main():
    sd = os.environ.get("SHIP_DIR", "/srv/homestyle/shipments")
    out = sys.argv[1]
    rows = json.load(open(os.path.join(sd, "rows.json"), encoding="utf-8"))
    try:
        alerts = json.load(open(os.path.join(sd, "alerts.json"), encoding="utf-8"))
    except OSError:
        alerts = {}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheets = {}
    for r in rows:
        sheets.setdefault(r.get("sheet") or "גיליון", []).append(r)
    head_fill = PatternFill("solid", fgColor="2F5D8A")
    for name, rs in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        ws.sheet_view.rightToLeft = True
        ws.append([h for h, _ in COLS] + ["התראות"])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center")
        for r in rs:
            line = []
            for _, k in COLS:
                v = r.get(k)
                if k in DATES:
                    try:
                        v = dt.date.fromisoformat(v) if v else None
                    except ValueError:
                        v = None
                line.append(v if v != "" else None)
            line.append(", ".join(alerts.get(r["id"], [])) or None)
            ws.append(line)
        for i, (_, k) in enumerate(COLS, start=1):
            if k in DATES:
                for cell in ws[get_column_letter(i)][1:]:
                    cell.number_format = "DD/MM/YY"
        widths = [10, 8, 10, 14, 34, 14, 6, 11, 11, 11, 12, 14, 22, 36, 14, 12, 28]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    tmp = out + ".tmp"
    wb.save(tmp)
    os.chmod(tmp, 0o640)
    os.replace(tmp, out)
    print("xlsx %d sheets, %d rows" % (len(sheets), len(rows)))


if __name__ == "__main__":
    main()
